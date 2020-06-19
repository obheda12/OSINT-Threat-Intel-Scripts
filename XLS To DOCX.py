import openpyxl,os,docx,re

os.chdir(r'C:\PATH')

wb = openpyxl.load_workbook('example.xlsx')
sheet = wb.get_sheet_by_name('Sheet1')
doc = docx.Document('example.docx')
fillObj = ('AAAA','BBBB','CCCC','DDDD')

for a in range(1,61):
    for b in range(1,5):
        fill = sheet.cell(row=a,column=b).value
        for x in range(len(fillObj)):
            inputRegex = re.compile(fillObj[x])
            inputRegex.sub(fill,doc)

        doc.save('example_' + fill + '.docx')
